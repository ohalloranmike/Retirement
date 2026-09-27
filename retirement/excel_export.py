from __future__ import annotations

from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from retirement.models import RetirementInputs, WithdrawalOrder
from retirement.projection import run_projection

INPUT_LABELS: list[tuple[str, str, str]] = [
    ("birth_year", "Birth year", "1965"),
    ("retirement_age", "Retirement age", "65"),
    ("life_expectancy_age", "Life expectancy (age)", "90"),
    ("planning_start_year", "Planning start year", "2025"),
    ("balance_401k", "401(k) balance ($)", "450000"),
    ("balance_traditional_ira", "Traditional IRA ($)", "120000"),
    ("balance_roth_ira", "Roth IRA ($)", "80000"),
    ("balance_taxable", "Taxable account ($)", "150000"),
    ("annual_salary", "Annual salary ($)", "120000"),
    ("employee_401k_contribution", "Employee 401(k) deferral ($/yr)", "15000"),
    ("employer_match_rate", "Employer match rate (0–1)", "0.5"),
    ("employer_match_up_to_pct_of_salary", "Match applies to first % of salary", "0.06"),
    ("annual_ira_contribution", "IRA contribution ($/yr)", "7000"),
    ("ira_is_roth", "IRA is Roth? (TRUE/FALSE)", "TRUE"),
    ("annual_return_pre_retirement", "Return pre-retirement (0–1)", "0.07"),
    ("annual_return_post_retirement", "Return in retirement (0–1)", "0.05"),
    ("pension_monthly_at_start", "Pension monthly at start ($)", "2500"),
    ("pension_start_age", "Pension start age", "65"),
    ("pension_cola_pct", "Pension COLA (0–1)", "0"),
    ("ss_monthly_at_fra", "SS monthly at FRA ($)", "3200"),
    ("ss_claim_age", "SS claim age (62–70)", "67"),
    ("ss_cola_pct", "SS COLA (0–1)", "0.02"),
    ("ss_fra_age", "SS full retirement age", "67"),
    ("spouse_ss_monthly_at_fra", "Spouse SS monthly at FRA ($)", "0"),
    ("spouse_ss_claim_age", "Spouse SS claim age", "67"),
    ("annual_spending_goal_today", "Spending goal (today's $/yr)", "85000"),
    ("inflation_pct", "Inflation (0–1)", "0.03"),
    ("withdrawal_order", "Withdrawal order", "taxable_traditional_roth"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(ws, row: int, ncol: int) -> None:
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def create_workbook(path: Path, inputs: RetirementInputs | None = None) -> Path:
    inputs = inputs or RetirementInputs(
        birth_year=1965,
        retirement_age=65,
        life_expectancy_age=90,
    )
    path = Path(path)
    wb = Workbook()
    ws_in = wb.active
    ws_in.title = "Inputs"
    ws_in["A1"] = "Field"
    ws_in["B1"] = "Description"
    ws_in["C1"] = "Value"
    _style_header(ws_in, 1, 3)

    field_values = {k: getattr(inputs, k) for k in inputs.__dataclass_fields__}
    for i, (key, label, default) in enumerate(INPUT_LABELS, start=2):
        ws_in.cell(row=i, column=1, value=key)
        ws_in.cell(row=i, column=2, value=label)
        val = field_values.get(key, default)
        if key == "withdrawal_order" and isinstance(val, WithdrawalOrder):
            val = val.value
        if key == "ira_is_roth":
            val = "TRUE" if val else "FALSE"
        ws_in.cell(row=i, column=3, value=val)

    ws_in.column_dimensions["A"].width = 32
    ws_in.column_dimensions["B"].width = 36
    ws_in.column_dimensions["C"].width = 18

    dv = DataValidation(
        type="list",
        formula1='"taxable_traditional_roth,traditional_taxable_roth,proportional"',
        allow_blank=False,
    )
    wo_row = 2 + next(i for i, (k, _, _) in enumerate(INPUT_LABELS) if k == "withdrawal_order")
    dv.add(ws_in[f"C{wo_row}"])
    ws_in.add_data_validation(dv)

    ws_help = wb.create_sheet("Instructions")
    ws_help["A1"] = "Retirement Planner (v1)"
    ws_help["A3"] = (
        "This workbook was exported from Retirement Planner (desktop GUI or Streamlit).\n"
        "1. Inputs — snapshot of assumptions at export time.\n"
        "2. Projection — year-by-year results computed in Python (not Excel formulas).\n"
        "3. To change assumptions, edit them in the app and export Excel again.\n"
        "4. Not tax advice; SS/pension are simplified models."
    )
    ws_help["A3"].alignment = Alignment(wrap_text=True)
    ws_help.column_dimensions["A"].width = 90

    _write_projection_sheet(wb, run_projection(inputs))

    wb.save(path)
    return path


def _write_projection_sheet(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Projection")
    cols = list(df.columns)
    for c, name in enumerate(cols, start=1):
        ws.cell(row=1, column=c, value=name)
    _style_header(ws, 1, len(cols))

    for r, row in enumerate(df.itertuples(index=False), start=2):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            if cols[c - 1] in ("return_rate",) and isinstance(val, (int, float)):
                cell.number_format = "0.00%"

    for c in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(c)].width = max(12, min(18, len(cols[c - 1]) + 2))

    # Income by source chart (retirement years only)
    ret_df = df[df["phase"] == "retirement"]
    if len(ret_df) >= 2:
        chart_start = len(df) + 3
        ws.cell(row=chart_start, column=1, value="Chart data (retirement)")
        headers = ["year", "pension_income", "ss_income", "withdrawal_total"]
        for i, h in enumerate(headers, start=1):
            ws.cell(row=chart_start + 1, column=i, value=h)
        for ri, row in enumerate(ret_df[headers].itertuples(index=False), start=chart_start + 2):
            for ci, val in enumerate(row, start=1):
                ws.cell(row=ri, column=ci, value=val)

        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Income by source (retirement)"
        chart.y_axis.title = "Dollars"
        data = Reference(
            ws,
            min_col=2,
            min_row=chart_start + 1,
            max_col=4,
            max_row=chart_start + 1 + len(ret_df),
        )
        cats = Reference(
            ws,
            min_col=1,
            min_row=chart_start + 2,
            max_row=chart_start + 1 + len(ret_df),
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 12
        chart.width = 24
        ws.add_chart(chart, f"F{chart_start}")

        bal_chart = LineChart()
        bal_chart.title = "Total portfolio balance"
        bal_data = Reference(ws, min_col=cols.index("balance_total") + 1, min_row=1, max_row=1 + len(df))
        bal_cats = Reference(ws, min_col=1, min_row=2, max_row=1 + len(df))
        bal_chart.add_data(bal_data, titles_from_data=True)
        bal_chart.set_categories(bal_cats)
        bal_chart.height = 12
        bal_chart.width = 24
        ws.add_chart(bal_chart, f"F{chart_start + 18}")
